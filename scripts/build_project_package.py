from pathlib import Path
import zipfile
ROOT=Path(__file__).resolve().parents[1]
TITLE="Implementation of Binary to Gray Code Converter in Quantum Dot Cellular Automata"
def read(p):
    f=ROOT/p
    return f.read_text(encoding="utf-8") if f.exists() else ""
summary=read(Path("docs/PROJECT_EXPLANATION.md"))
workflow=read(Path("docs/WORKFLOW.md"))
structure=read(Path("docs/PROJECT_STRUCTURE.md"))
source=read(Path("docs/SOURCE_MATERIALS.md"))
truth=[(format(b,"04b"),format(b^(b>>1),"04b")) for b in range(16)]
lines=["# "+TITLE,"","## Project Identification","- Institution: CMR Institute of Technology, Hyderabad","- Department: Electronics and Communication Engineering","- Academic Year: 2024-25","- Guide: Dr. G. Rajender","- Members: Aleti Harneeth Reddy; Alladi Nithin Kumar; GVN Bharadwaj; Kavali Harshavardhan","","## Abstract","This project implements a Binary-to-Gray Code Converter using Quantum Dot Cellular Automata (QCA) design concepts.","","## Objectives","1. Study QCA as a nanoscale digital design technology.","2. Implement Binary-to-Gray code conversion logic.","3. Provide Verilog reference modules and testbenches.","","## Boolean Equations","### 2-bit","G1 = B1","G0 = B1 XOR B0","","### 4-bit","G3 = B3","G2 = B3 XOR B2","G1 = B2 XOR B1","G0 = B1 XOR B0","","## 4-bit Truth Table","| Binary | Gray |","|---|---|"]
lines += ["| "+b+" | "+g+" |" for b,g in truth]
lines += ["","## QCA Design Concepts","- QCA cells represent binary information through electron configuration.","- QCA wires transfer cell polarization.","- Majority gates and inverters are fundamental QCA logic primitives.","- Clocking controls computation and signal propagation.","","## Workflow",workflow,"","## Project Structure",structure,"","## Source Materials",source,"","## Verilog Files","- verilog/xor_gate.v","- verilog/xor_gate_tb.v","- verilog/binary_to_gray_2bit.v","- verilog/binary_to_gray_2bit_tb.v","- verilog/binary_to_gray_4bit.v","- verilog/binary_to_gray_4bit_tb.v","","## Verification","Gray = Binary XOR (Binary >> 1). The included testbenches exercise the reference logic.","","## Tools Mentioned","The source report references QCA Designer and Microwind Lite.","","## Deliverables","Generated PDF, DOCX, PPTX, XLSX and ZIP files are built from the committed project documentation and implementation files."]
report="\n".join(lines)
(ROOT/"docs/COMPLETE_PROJECT_REPORT.md").write_text(report,encoding="utf-8")

from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
styles=getSampleStyleSheet()
story=[Paragraph(TITLE,styles["Title"]),Paragraph("Generated complete project report.",styles["Normal"]),Spacer(1,12)]
for line in lines:
    if line.startswith("# "): story.append(Paragraph(line[2:],styles["Title"]))
    elif line.startswith("## "): story.append(Paragraph(line[3:],styles["Heading1"]))
    elif line.startswith("### "): story.append(Paragraph(line[4:],styles["Heading2"]))
    elif line.startswith("- "): story.append(Paragraph("• "+line[2:],styles["BodyText"]))
    elif line.strip() and not line.startswith("|"): story.append(Paragraph(line,styles["BodyText"])); story.append(Spacer(1,3))
story.append(Paragraph("Truth Table",styles["Heading1"]))
table=Table([["Binary","Gray"]]+truth,colWidths=[90,90])
table.setStyle(TableStyle([("GRID",(0,0),(-1,-1),0.5,colors.black),("BACKGROUND",(0,0),(-1,0),colors.lightgrey),("ALIGN",(0,0),(-1,-1),"CENTER")]))
story.append(table)
(ROOT/"reports").mkdir(exist_ok=True)
SimpleDocTemplate(str(ROOT/"reports/project_report.pdf"),pagesize=A4,rightMargin=42,leftMargin=42,topMargin=42,bottomMargin=42).build(story)

from docx import Document
d=Document(); d.add_heading(TITLE,0); d.add_paragraph("Generated project notes and implementation reference.")
for name,txt in [("Project Summary",summary),("Workflow",workflow),("Structure",structure),("Source Materials",source)]:
    d.add_heading(name,1); d.add_paragraph(txt)
d.add_heading("Truth Table",1)
tb=d.add_table(rows=1,cols=2); tb.rows[0].cells[0].text="Binary"; tb.rows[0].cells[1].text="Gray"
for b,g in truth:
    c=tb.add_row().cells; c[0].text=b; c[1].text=g
d.save(ROOT/"reports/project_notes.docx")

d2=Document(); d2.add_heading("Project Acknowledgement and Verilog Reference",0); d2.add_paragraph("Generated repository copy for the Binary-to-Gray QCA project.")
d2.add_heading("Included modules",1)
for p in ["verilog/xor_gate.v","verilog/binary_to_gray_2bit.v","verilog/binary_to_gray_4bit.v","verilog/xor_gate_tb.v","verilog/binary_to_gray_2bit_tb.v","verilog/binary_to_gray_4bit_tb.v"]: d2.add_paragraph(p,style="List Bullet")
d2.save(ROOT/"reports/bcd_to_gcd_acknowledgement.docx")

from openpyxl import Workbook
wb=Workbook(); ws=wb.active; ws.title="Project Summary"
for row in [["Project Title",TITLE],["Institution","CMR Institute of Technology, Hyderabad"],["Department","Electronics and Communication Engineering"],["Academic Year","2024-25"],["Guide","Dr. G. Rajender"],["Members","Aleti Harneeth Reddy; Alladi Nithin Kumar; GVN Bharadwaj; Kavali Harshavardhan"],["Core Formula","Gray = Binary XOR (Binary >> 1)"]]: ws.append(row)
ws2=wb.create_sheet("4-bit Truth Table"); ws2.append(["Binary","Gray"])
for b,g in truth: ws2.append([b,g])
for s in wb.worksheets: s.column_dimensions["A"].width=32; s.column_dimensions["B"].width=65
(ROOT/"data").mkdir(exist_ok=True); wb.save(ROOT/"data/Project_Details_and_Truth_Table.xlsx")

from pptx import Presentation
prs=Presentation()
slides=[("Title",[TITLE,"CMR Institute of Technology - ECE - 2024-25"]),("Objectives",["Study QCA concepts","Implement Binary-to-Gray conversion","Provide Verilog reference and verification"]),("4-bit Equations",["G3 = B3","G2 = B3 XOR B2","G1 = B2 XOR B1","G0 = B1 XOR B0"]),("QCA Concepts",["QCA cells and polarization","Majority gates and inverters","Clock zones and signal propagation"]),("Workflow",["Specification -> equations -> Verilog -> QCA concepts -> verification -> documentation"]),("Verification",["16 possible 4-bit inputs","Gray = Binary XOR (Binary >> 1)","Testbenches included"]),("Deliverables",["PDF","DOCX","PPTX","XLSX","Verilog","ZIP"])]
for title,bullets in slides:
    slide=prs.slides.add_slide(prs.slide_layouts[0] if title=="Title" else prs.slide_layouts[1]); slide.shapes.title.text=title
    if title=="Title": slide.placeholders[1].text="\n".join(bullets)
    else:
        tf=slide.placeholders[1].text_frame; tf.clear()
        for i,b in enumerate(bullets):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=b
prs.save(ROOT/"presentation/mini_project_presentation.pptx")

dist=ROOT/"dist"; dist.mkdir(exist_ok=True); zpath=dist/"Binary_to_Gray_QCA_Complete_Project.zip"
with zipfile.ZipFile(zpath,"w",zipfile.ZIP_DEFLATED) as z:
    for p in ROOT.rglob("*"):
        if p.is_file() and ".git" not in p.parts and p != zpath and p.name!="journal_paper_womens_safety.docx": z.write(p,p.relative_to(ROOT))
print(zpath)
